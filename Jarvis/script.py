import os
from pathlib import Path

from docx import Document
from dotenv import load_dotenv
from google import genai
from openai import OpenAI

load_dotenv("dev.env")

def gemini():
    client = genai.Client(api_key=os.environ["GEMINI_API_KEY"])
    response = client.models.generate_content(
        model="gemini-3.6-flash",
        contents="Say hi in one word."
    )
    print(response.text)


# GitHub Models retired July 30, 2026 — this no longer works, kept for reference only
def github():
    client = OpenAI(
        base_url="https://models.github.ai/inference",
        api_key=os.environ["GITHUB_TOKEN"],
    )
    response = client.chat.completions.create(
        model="openai/gpt-4o-mini",
        messages=[{"role": "user", "content": "Say hi in one word."}],
    )
    print(response.choices[0].message.content)

def docx_reader(path):
    document = Document(path)
    print(document)

    print("=== PARAGRAPHS ===")
    for i, para in enumerate(document.paragraphs):
        print(i, repr(para.text))

    print("=== TABLES ===")
    for t, table in enumerate(document.tables):
        print(f"--- table {t} ---")
        for row in table.rows:
            cells = [cell.text for cell in row.cells]
            print(cells)

if __name__ == "__main__":
    docx_reader(Path(__file__).resolve().parent.parent / ".agents/workspace/Rohin_Agrawal_Resume.docx")
    docx_reader(Path(__file__).resolve().parent.parent / ".agents/workspace/Biodata.docx")