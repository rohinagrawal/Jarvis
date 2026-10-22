from docx import Document


def docx_reader(path):
    document = None
    with open(path, "r") as file:
        document = Document(file)
        print(document)

    print("=== PARAGRAPHS ===")
    for i, para in enumerate(document.paragraphs):
        print(i, repr(para.text))

    print("=== TABLES ===")
    for t, table in enumerate(document.tables):
        print(f"--- table {t} ---")
        for row in table.rows:
            cells = [cell.text for cell in row.cells]
            print(cells)

    return document

def file_reader(path):
    content = None
    with open(path, "r", encoding="utf-8") as file:
        content = file.read()
        # print(content)
    return content

def count_token_approx(text:str) -> int:
    # Count the no of words in a text to find approximate tokens
    return len(text.split())